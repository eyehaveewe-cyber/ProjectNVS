"""Archived builder for the original version 0.2 woodland design addendum.

Run with the Codex primary Python runtime. Implementation verification is
recorded separately by the game's replayable tests and PROGRESS.md. This
historical builder refuses to overwrite an existing design document. Use
build-map-expansion-notes.py to update the current addendum in place.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ProjectNVS_Woodland_Expansion.docx"
if OUTPUT.exists():
    raise SystemExit(
        "The current design addendum already exists. This archived 0.2 builder "
        "will not overwrite its edits. Use tools/build-map-expansion-notes.py."
    )
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.66)
section.left_margin = section.right_margin = Inches(0.72)
section.footer_distance = Inches(0.3)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for name, size in [("Title", 26), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 12)]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10 if "Heading" in name else 0)
    style.paragraph_format.space_after = Pt(6)

doc.core_properties.title = "Project NVS woodland expansion"
doc.core_properties.subject = "Design addendum and implementation scope"
doc.core_properties.author = "Channel Hideo design assistant"
doc.core_properties.version = "0.2"

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Project NVS  |  23 October 2026  |  ")
run.font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def p(text, style=None):
    return doc.add_paragraph(text, style=style)


def h(text, level=1):
    doc.add_heading(text, level=level)


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, header in enumerate(headers):
        t.rows[0].cells[i].text = header
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    tblpr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    tblpr.append(borders)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "85")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tcpr.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "294A43" if ri == 0 else "F2F5F4" if ri % 2 == 0 else "FFFFFF")
            tcpr.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.03
                for run in paragraph.runs:
                    run.font.size = Pt(9.3)
                    if ri == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    p("").paragraph_format.space_after = Pt(0)
    return t


doc.add_heading("Project NVS woodland expansion", 0)
p("Design addendum 0.2  |  23 October 2026  |  Director Jonathan Sanders", "Subtitle")
p("The first expansion should let the faun wander and make music beyond the original restoration prompts. A larger woodland gives quiet travel room to matter; a squirrel, recovering wildlife and returning dancers give that travel occasional purpose. The human direction remains staged development, with a playable first pass before a larger mythological story.")

h("Human direction and authorization")
p("The entries below preserve the origin of the ideas. H001 to H004 come from the Project NVS section of the supplied playtest notes; H005 and H006 come from the current request. The instruction to implement changes authorizes this pass. The exact rules and placements on the next page are assistant implementation choices, rather than additional claims about the director's original design.")
table(["ID", "Human contribution", "Scope and priority"], [
    ("H001", "Play MBean1's fuller keyboard freely across multiple octaves.", "First pass; explicit current request."),
    ("H002", "Listen and answer melodies; keep improvisation available. A note highway is an option.", "Listening retained; highway deferred."),
    ("H003", "Cleanse polluted animals and grow sprouts or vines from damaged woodland.", "Two distinct recovery examples in this pass."),
    ("H004", "Expansive quiet woodland, small creatures and discoveries unrelated to pollution.", "First pass; explicit current request."),
    ("H005", "Chase a squirrel to unlock a useful song, such as rain; add occasional collecting.", "First pass; song effect and optional collecting."),
    ("H006", "More animation on clearing corruption, dancing creatures, dryads, Artemis as a woman or a giant deer.", "Celebrations and giant-deer encounter first; human form open."),
], [0.5, 3.92, 2.54])

h("Experience to protect")
p("Music visibly helps living things. Exploration can include stretches with little happening. Collecting should invite a detour without converting every path into a checklist, and a missed phrase should allow an immediate retry. The existing nine-sanctuary journey and heart-tree ending remain the foundation.")

doc.add_page_break()
h("First playable milestone")
p("These assistant choices implement H001 and H003 to H006 in the native Godot project. The reports on page three distinguish source behavior, exercised interactions and remaining playtest limits. Timings and placements are assistant tuning choices that can change.")

h("An instrument that remains usable", 2)
p("The actual MBean1 layout is playable with held polyphony and a bank of 61 PCM notes. A S D F J K L play the current octave; Q W E R U I O play one higher; Z X C V M comma period play one lower. T Y and B N provide outer notes. Tab raises or lowers the pipes, Space sharpens eligible notes, Page Up and Page Down change octave, and Ctrl or Shift provide MBean octave modifiers. Pause, focus loss, mute and lowering the pipes release held voices. Sanctuary hints follow the keyboard labels. [H001, H002]")

h("Quiet routes with occasional reasons to stop", 2)
p("The walkable square expands from 52.4 by 52.4 to 150 by 150 units, retaining the central valley. That is about 8.2 times the bounded surface, rather than a measure of content. Six connected routes total 557.65 units and link six clearings, including Moonmere, the dryad glade, reedwater bend, stump nursery and moon meadow. Trees, signs and trail surfaces guide travel; 18 optional seeds reward detours without gating sanctuary progress. [H004, H005]")

h("The squirrel and its rain song", 2)
p("A squirrel on the northwestern trail leads through six waypoints toward Moonmere, waiting when the player falls behind. Completing the pursuit unlocks Raincall; playing its phrase calls a 45-second shower. Rain near the nursery starts 14 seconds of small-sprout growth. The song unlock persists through save and Continue. Wrong phrases and an abandoned pursuit leave another opportunity. [H005]")

h("Different things can recover differently", 2)
p("The oily fox at reedwater bend has its own cleansing phrase. Its coating disappears and it begins moving freely. Rain encourages tiny nursery plants; healed scorched sites grow young shoots over five seconds, leaving evidence of the old damage. Clearing a sanctuary triggers light rings, petals, hopping animals and a dancing dryad. Some woodland residents remain afterward. [H003, H006]")

h("A first meeting with Artemis", 2)
p("A large luminous deer appears in the moon meadow after three sanctuaries recover. Its antlers and scale distinguish Artemis from ordinary wildlife; a short encounter references a silent forest. The deer is the first prototype of the director's suggested forms. The woman form and the larger cause of the damage remain open, with no adopted antagonist or campaign plot. [H006]")

doc.add_page_break()
h("Coverage and proof of completion")
p("Godot 4.7.2 reports pass 39 core, 83 music, 10 route, 45 woodland integration and 14 controller-traversal checks. Captures show trails, rain, celebration and Artemis. Integration relocates the faun between discoveries; physical squirrel travel and stated keyboard phrases use game ticks and input events. Traversal has separate evidence below.")
table(["Report", "Recorded evidence and limit"], [
    ("music-keyboard", "83 checks: mapping, held notes, releases and cleanup. Core integration checks all nine sanctuary phrases and the ending."),
    ("woodland-routes", "10 checks and 1,359 clear sampled points: trail topology, boundary passages, seeds and clearings. Sampling alone does not prove player traversal."),
    ("woodland", "45 checks: physical squirrel travel, Raincall gating, timed rain and growth, fox recovery, cast, dryad dialogue, Artemis and save migration."),
    ("woodland-traversal", "14 checks: real controller follows 557.49 units through 33 waypoints in 3,489 steps. Each of six routes starts with relocation; 174.45 simulated seconds is not a frame-rate measurement."),
    ("Native captures", "Godot captures show trails, rain, celebration and Artemis. They establish appearance, not musical enjoyment or frame rate."),
    ("godot core", "39 checks: scene, walking, restoration, heart ending, save and Continue, malformed input, rest and PCM. Earlier browser reports cover the legacy edition."),
], [1.44, 5.52])

h("Remaining questions and deferred work")
p("The note highway is deferred; this pass keeps free timing and listening hints. Artemis's human form, the wider conflict and an absent familiar melody as a story event remain open. A fuller music-learning curriculum and deeper daily care are outside this milestone, without implying rejection. [H002, H004, H006]")
p("Check for instrument hotkey collisions, indistinguishable quiet routes, frustrating pursuit and celebrations that hide the faun. Test both input modes and return paths, let the squirrel wait, and inspect recovery from the normal camera before adding more effects.")

h("Recommended next decision")
p("After playing, assess whether rain and creature care reward wandering. The assistant recommendation is to develop one more musical relationship before adding a region. Separately decide Artemis's encounter frequency and whether both forms belong to one character. These story choices remain unadopted.")

h("Sources and revision history")
p("S1  Crunch_Time_Jam_Playtest_Notes_2026-10-22(2).docx, section 03 Project NVS. The Foxlike and Project Klax sections do not inform this addendum.")
p("S2  Jonathan Sanders's current request, 23 October 2026, including staged development. H005 and the new celebration and form details in H006 originate here.")
p("S3  https://github.com/eyehaveewe-cyber/ProjectNVS. Keyboard source: MBean1-Astra, piano-performance.js. The implementation record holds final source links and test results.")
p("Revision 0.2  Records implemented rules, controls and native verification reports. Publication is pending. Complements the earlier Channel Hideo review; human authorship and open story decisions remain distinct.")

for element in doc._element.iter():
    for border in list(element.findall(qn("w:pBdr"))):
        element.remove(border)
for style in doc.styles:
    for border in list(style._element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

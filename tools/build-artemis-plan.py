"""Append the reviewed Terra handoff while preserving earlier design pages."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'docs/ProjectNVS_Woodland_Expansion.docx'
TITLE = 'Artemis expansion plan for Terra'
doc = Document(OUTPUT)
if float(doc.core_properties.version or '0') > 0.5:
    raise SystemExit('A newer design revision exists. Preserve its edits.')
for paragraph in list(doc.paragraphs):
    if paragraph.text == TITLE:
        element = paragraph._p
        while element is not None and element.tag != qn('w:sectPr'):
            following = element.getnext()
            element.getparent().remove(element)
            element = following
        break
for paragraph in doc.paragraphs:
    if paragraph.style.name == 'Subtitle' and paragraph.text.startswith('Design addendum '):
        paragraph.text = 'Design addendum 0.5  |  28 October 2026  |  Director Jonathan Sanders'
heading = doc.add_heading(TITLE, 1)
heading.paragraph_format.page_break_before = True
for line in (ROOT / 'docs/TERRA_EXPANSION_PLAN.txt').read_text().splitlines()[1:]:
    if not line.strip():
        continue
    if line.isupper():
        doc.add_heading(line.capitalize(), 2)
    else:
        doc.add_paragraph(line)
doc.core_properties.version = '0.5'
doc.save(OUTPUT)
print(OUTPUT)

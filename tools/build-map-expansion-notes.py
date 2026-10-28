"""Update the existing woodland addendum without rebuilding its prior pages.

Run with the Codex primary Python runtime after the native verification scripts.
All earlier human text and the original contribution tables remain in the DOCX.
Only this tool's current appendix is replaced when rerun.
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ProjectNVS_Woodland_Expansion.docx"
APPENDIX = "Far woodland map expansion"


def report(name):
    value = json.loads((ROOT / "docs" / name).read_text())
    assert value["status"] == "passed" and not value["failures"], name
    return value


routes = report("woodland-routes-verification.json")
traversal = report("woodland-traversal-verification.json")
journey = report("woodland-verification.json")
assert routes["routes"] == 11 and routes["bounds"] == 125.0
assert traversal["authored_distance"] > 1600.0
assert any("continue" in check.lower() and "outer" in check.lower() for check in journey["checks"])

doc = Document(OUTPUT)
for paragraph in list(doc.paragraphs):
    if paragraph.text in [APPENDIX, "Map expansion on 27 October 2026"]:
        element = paragraph._p
        while element is not None and element.tag != qn("w:sectPr"):
            following = element.getnext()
            element.getparent().remove(element)
            element = following
        break

patches = {
    "Design addendum 0.2  |  23 October 2026  |  Director Jonathan Sanders":
        "Design addendum 0.3  |  28 October 2026  |  Director Jonathan Sanders",
    "Design addendum 0.3  |  27 October 2026  |  Director Jonathan Sanders":
        "Design addendum 0.3  |  28 October 2026  |  Director Jonathan Sanders",
    "First playable milestone": "First woodland pass from 23 October 2026",
    "Coverage and proof of completion": "Verification from 23 October 2026",
    "Recommended next decision": "Recommendation recorded on 23 October 2026",
    "Godot 4.7.2 reports pass 39 core, 83 music, 10 route, 45 woodland integration and 14 controller-traversal checks. Captures show trails, rain, celebration and Artemis. Integration relocates the faun between discoveries; physical squirrel travel and stated keyboard phrases use game ticks and input events. Traversal has separate evidence below.":
        "The recovered 23 October build recorded 39 core, 83 music, 10 route, 45 woodland and 14 controller-traversal checks. Its captures show trails, rain, celebration and Artemis. Integration relocated the faun between discoveries; squirrel travel and keyboard phrases used game ticks and input events. The expanded map has a current verification record on the next page.",
}
for paragraph in doc.paragraphs:
    if paragraph.text in patches:
        paragraph.text = patches[paragraph.text]
for section in doc.sections:
    for paragraph in section.footer.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace("23 October 2026", "28 October 2026").replace("27 October 2026", "28 October 2026")
doc.core_properties.version = "0.3"


def p(text):
    return doc.add_paragraph(text)


def h(text, level=2):
    return doc.add_heading(text, level=level)


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
        t.rows[0].cells[i].text = headers[i]
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), val)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "85")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "39443F" if ri == 0 else "F2F5F4" if ri % 2 == 0 else "FFFFFF")
            properties.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.03
                for run in paragraph.runs:
                    run.font.size = Pt(9.3)
                    if ri == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    return t


heading = h(APPENDIX, 1)
heading.paragraph_format.page_break_before = True
p("H007 is the director's approved request to continue the recovered build and expand the map. The wording was \"let's continue where you left off and do an expansion to map size from there\". This pass implements that request in native Godot. The dimensions, ring layout, new place names and prop arrangements below are AI implementation choices within H007, building on H004 and H005.")

table(["Measure", "Recovered woodland", "Expanded woodland"], [
    ("Walkable square", "150 by 150 units", "250 by 250 units"),
    ("Bounded surface", "22,500 square units", "62,500 square units, 2.78 times prior"),
    ("Authored trails", f"6 routes, {routes['original_trail_length']:.2f} units", f"11 routes, {routes['trail_length']:.2f} units"),
    ("Named clearings", "6", "11"),
    ("Optional wild seeds", "18", "30, with original indices retained"),
], [1.56, 2.20, 3.20])

h("Routes and a reason to wander")
p("Five added trail sections form an outer circuit with returns through Moonmere, the Listening Stones, Reedwater, Artemis' Meadow and the stump nursery. Fern Hollow has a fallen log and ferns; Wind Stones has tall stone silhouettes; Willow Pool has a drooping tree; Sunwash Meadow keeps an open flowered center; the Elder Ring surrounds a quiet clearing. Twelve scattered seeds offer optional detours.")
p("For example, leave Moonmere for Fern Hollow, continue toward Wind Stones, then return by the older northern walk or follow the circuit east. No new song or restoration is required. The journal now has a labeled full woodland map and route guide, and the small map covers the wider country. Native version-two saves retain the older discoveries and can resume beyond the former boundary.")

h("Current verification and limits")
p(f"Godot {routes['engine']} passes {len(routes['checks'])} route checks across {routes['all_samples']:,} collision samples, with no blocked samples. The controller report passes {len(traversal['checks'])} checks and walks the routes forward and backward with only a starting relocation per route. Woodland integration passes {len(journey['checks'])} checks, including actual outer-country save and Continue, old and added seed IDs, and legacy saves. Reports are woodland-routes-verification.json, woodland-traversal-verification.json and woodland-verification.json in docs.")
p("Native GL Compatibility renders of the five added landmarks and the labeled journal map were inspected. The Elder Ring was revised to open its foreground view after crowns initially hid the clearing. These captures use software graphics and establish appearance.")
p("These checks establish connectivity, controller traversal, optional discoveries and persistence. They do not establish first-time navigation, musical enjoyment or target-hardware performance. The note highway, Artemis' human form and the wider mythological conflict remain deferred. The next playtest should assess whether the signed loops are readable and whether the longer quiet stretches invite wandering.")
p("S4  Director request, 27 October 2026. Revision 0.3 completed on 28 October 2026 preserves the earlier design and evidence, records H007 and the current map implementation, and keeps publication separate from verification. The source baseline was recovered from ProjectNVS_Woodland_Source_74fb1e5.zip.")

doc.save(OUTPUT)
print(OUTPUT)
